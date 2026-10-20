#!/usr/bin/env python3
"""Regression tests for the established Step2 platform-mapping condition."""

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from update_erp_step2 import (
    PRODUCT_SUBJECTS,
    REQUIRED_MAIN,
    append_tracking_skus,
    headers,
    normalize_product_subjects,
    platform_sales_eligible_rows,
    prepare_tracking_data,
    product_subject,
    short_product_subject_cells,
)


def build_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品表"
    sheet.append([])
    sheet.append(REQUIRED_MAIN)
    return workbook, sheet, headers(sheet)


def main():
    workbook, sheet, columns = build_workbook()
    sheet.cell(3, columns["平台SKU"]).value = "EXISTING-MSKU"
    sheet.cell(3, columns["仓库SKU"]).value = "EXISTING-SKU"

    thin = Side(style="thin", color="000000")
    for row in (4, 5):
        for column in range(1, len(REQUIRED_MAIN) + 1):
            sheet.cell(row, column).border = Border(
                left=thin, right=thin, top=thin, bottom=thin
            )

    tracking = pd.DataFrame([
        {
            "平台SKU": f"NEW-MSKU-{index}",
            "SKU": f"NEW-SKU-{index}",
            "预计开售时间": datetime(2026, 10, 14),
        }
        for index in range(1, 4)
    ])
    tracking.loc[0, "SKU"] = "EXISTING-SKU"
    appended, _ = append_tracking_skus(
        workbook, sheet, columns, tracking, legacy_name=None
    )
    assert appended == 3
    assert sheet.max_row == 6
    assert [sheet.cell(row, columns["平台SKU"]).value for row in range(4, 7)] == [
        "NEW-MSKU-1",
        "NEW-MSKU-2",
        "NEW-MSKU-3",
    ]
    assert platform_sales_eligible_rows(sheet, columns) == [3, 4, 5, 6]

    raw_tracking = pd.DataFrame([
        {"平台SKU": "KEEP-BLANK-GROUP", "SKU": "SKU-1", "主体": "KM一部一组", "小组": "", "预计开售时间": "2026/10", "新品期截止月": "2026/12", "预测销量完成比例": "10%"},
        {"平台SKU": "KEEP-BLANK-GROUP", "SKU": "SKU-2", "主体": "KM一部一组", "小组": "", "预计开售时间": "2026/11", "新品期截止月": "2027/01", "预测销量完成比例": "20%"},
        {"平台SKU": "", "SKU": "SKU-3", "主体": "KM一部三组", "小组": "", "预计开售时间": "2026/12", "新品期截止月": "2027/02", "预测销量完成比例": "0%"},
        {"平台SKU": "DROP-OTHER-SUBJECT", "SKU": "SKU-4", "主体": "其他主体", "小组": "KM一部一组", "预计开售时间": "2026/12", "新品期截止月": "2027/02", "预测销量完成比例": "0%"},
    ])
    prepared = prepare_tracking_data(raw_tracking)
    assert len(raw_tracking) == 4
    assert list(prepared["平台SKU"]) == ["KEEP-BLANK-GROUP"]
    assert prepared.iloc[0]["SKU"] == "SKU-1"

    assert product_subject("KM一部一组") == "KM一部一组【KM一部一组】"
    assert product_subject("KM一部二组【KM一部二组】") == "KM一部二组【KM一部二组】"
    assert product_subject("其他主体") == ""

    sheet.cell(3, columns["主体"]).value = "KM一部一组"
    sheet.cell(4, columns["主体"]).value = PRODUCT_SUBJECTS["KM一部二组"]
    sheet.cell(5, columns["主体"]).value = None
    sheet.cell(6, columns["主体"]).value = "其他主体"
    assert normalize_product_subjects(sheet, columns) == 1
    assert sheet.cell(3, columns["主体"]).value == PRODUCT_SUBJECTS["KM一部一组"]
    assert sheet.cell(4, columns["主体"]).value == PRODUCT_SUBJECTS["KM一部二组"]
    assert sheet.cell(5, columns["主体"]).value is None
    assert sheet.cell(6, columns["主体"]).value == "其他主体"
    assert short_product_subject_cells(sheet, columns) == []

    for row in (3, 4, 5, 6):
        sheet.cell(row, columns["店铺别名"]).value = "KM-US"
    assert platform_sales_eligible_rows(sheet, columns) == []

    sheet.cell(6, columns["店铺别名"]).value = None
    sheet.cell(6, columns["主体"]).value = PRODUCT_SUBJECTS["KM一部一组"]
    assert platform_sales_eligible_rows(sheet, columns) == []
    workbook.close()
    print("Step2 source-preservation and legacy-condition regression test passed.")


if __name__ == "__main__":
    main()
