#!/usr/bin/env python3
"""Apply ERP and Step 2 rules to a BI-refreshed KM operations workbook."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.formula.translate import Translator


ALLOWED_GROUPS = ("KM一部一组", "KM一部二组", "KM一部三组")
PRODUCT_SUBJECTS = {
    subject: f"{subject}【{subject}】" for subject in ALLOWED_GROUPS
}
LEGACY_NAMES = ("产品表-老三(含清货淘汰)", "产品表-老三（含清货淘汰）")
ERRORS = ("#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NUM!")
BASE_FIELDS = ("店铺别名", "主体", "运营", "备货状态")
REQUIRED_MAIN = [
    "店铺别名", "主体", "产品名称", "三级品类", "产品销售状态", "平台SKU",
    "仓库SKU", "ASIN", "品类等级", "运营", "备货状态", "新品期开售月",
    "新品期截止月", "最新新品成功率",
] + [f"标签{i:02d}" for i in range(1, 11)]


def clean(value):
    return "" if value is None else str(value).strip()


def normalized(value):
    return re.sub(r"\s+", "", clean(value))


def headers(sheet, row=2):
    return {normalized(cell.value): cell.column for cell in sheet[row] if normalized(cell.value)}


def require(mapping, names, label):
    missing = [name for name in names if normalized(name) not in mapping]
    if missing:
        raise ValueError(f"{label}缺少字段: {', '.join(missing)}")


def group_mask(series):
    return series.astype(str).str.contains("|".join(ALLOWED_GROUPS), regex=True, na=False)


def canonical_subject(value):
    text = clean(value)
    for subject in ALLOWED_GROUPS:
        pattern = rf"(?<![\u4e00-\u9fffA-Za-z0-9]){re.escape(subject)}(?![\u4e00-\u9fffA-Za-z0-9])"
        if re.search(pattern, text):
            return subject
    return ""


def product_subject(value):
    """Return the required long-form subject for 产品表 writes."""
    return PRODUCT_SUBJECTS.get(canonical_subject(value), "")


def normalize_product_subjects(sheet, cols):
    """Convert exact short subjects in 产品表 while preserving blanks and other values."""
    subject_column = cols[normalized("主体")]
    changed = 0
    for row in range(3, sheet.max_row + 1):
        cell = sheet.cell(row, subject_column)
        text = clean(cell.value)
        if text not in PRODUCT_SUBJECTS:
            continue
        cell.value = PRODUCT_SUBJECTS[text]
        changed += 1
    return changed


def short_product_subject_cells(sheet, cols):
    """Return 产品表 cells that still contain an exact short-form subject."""
    subject_column = cols[normalized("主体")]
    return [
        sheet.cell(row, subject_column).coordinate
        for row in range(3, sheet.max_row + 1)
        if clean(sheet.cell(row, subject_column).value) in PRODUCT_SUBJECTS
    ]


def subject_mask(series):
    return series.astype(str).str.strip().isin(ALLOWED_GROUPS)


def prepare_tracking_data(source):
    """Return a matching copy while leaving the complete export untouched."""
    require(
        {normalized(c): c for c in source.columns},
        ["平台SKU", "SKU", "主体", "预计开售时间", "新品期截止月", "预测销量完成比例"],
        "ERP新品成功率",
    )
    matching = source[subject_mask(source["主体"])].copy()
    matching = matching[matching["平台SKU"].astype(str).str.strip() != ""]
    return matching.drop_duplicates("平台SKU", keep="first")


def parse_date(value):
    if value is None or clean(value) == "":
        return None
    parsed = pd.to_datetime(value, errors="coerce")
    return None if pd.isna(parsed) else parsed.date()


def sheet_digest(sheet):
    payload = {
        "title": sheet.title,
        "state": sheet.sheet_state,
        "freeze": str(sheet.freeze_panes),
        "filter": str(sheet.auto_filter.ref),
        "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
        "rows": [],
        "row_dims": [],
        "col_dims": [],
    }
    for row in sheet.iter_rows():
        payload["rows"].append([
            (cell.coordinate, cell.value, cell.data_type, cell.style_id, cell.number_format)
            for cell in row if cell.value is not None or cell.has_style
        ])
    payload["row_dims"] = sorted((key, dim.height, dim.hidden, dim.outlineLevel) for key, dim in sheet.row_dimensions.items())
    payload["col_dims"] = sorted((key, dim.width, dim.hidden, dim.outline_level) for key, dim in sheet.column_dimensions.items())
    data = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
    return hashlib.sha256(data).hexdigest()


def replace_sheet(workbook, name, frame, header_row=1):
    if name in workbook.sheetnames:
        del workbook[name]
    sheet = workbook.create_sheet(name)
    if header_row == 2:
        sheet.append([None] * len(frame.columns))
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append([None if pd.isna(value) else value for value in row])


def formula_errors(workbook):
    found = set()
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value for token in ERRORS):
                    found.add(f"{worksheet.title}!{cell.coordinate}:{cell.value}")
    return found


def copy_row_style(sheet, source_row, target_row, max_col):
    for column in range(1, max_col + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy.copy(source.alignment)
        if source.protection:
            target.protection = copy.copy(source.protection)
        if source.data_type == "f" and isinstance(source.value, str):
            target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def last_content_row(sheet, first_data_row=3):
    """Return the last row containing a value or formula, ignoring styled blanks."""
    for row in range(sheet.max_row, first_data_row - 1, -1):
        if any(
            sheet.cell(row, column).value not in (None, "")
            for column in range(1, sheet.max_column + 1)
        ):
            return row
    return first_data_row - 1


def append_tracking_skus(workbook, sheet, cols, tracking, legacy_name):
    """Apply the established new-SKU append rule and return append/exclusion counts."""
    existing_platform_skus = set()
    for row in range(3, sheet.max_row + 1):
        platform_sku = clean(sheet.cell(row, cols[normalized("平台SKU")]).value)
        if platform_sku:
            existing_platform_skus.add(platform_sku)

    legacy_platform_skus = set()
    if legacy_name:
        legacy_cols = headers(workbook[legacy_name])
        column = legacy_cols.get(normalized("平台SKU"))
        if column:
            legacy_platform_skus.update(
                clean(workbook[legacy_name].cell(row, column).value)
                for row in range(3, workbook[legacy_name].max_row + 1)
            )

    exclusions = existing_platform_skus | legacy_platform_skus
    appended = 0
    append_start_row = last_content_row(sheet) + 1
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    for _, source in tracking.iterrows():
        platform_sku = clean(source["平台SKU"])
        warehouse_sku = clean(source["SKU"])
        launch = parse_date(source["预计开售时间"])
        if (
            not platform_sku
            or platform_sku in exclusions
            or launch is None
            or launch < datetime(2026, 1, 1).date()
        ):
            continue
        target_row = append_start_row + appended
        if target_row > sheet.max_row:
            copy_row_style(sheet, target_row - 1, target_row, 43)
        sheet.cell(target_row, cols[normalized("平台SKU")]).value = platform_sku
        sheet.cell(target_row, cols[normalized("仓库SKU")]).value = warehouse_sku or None
        sheet.cell(target_row, cols[normalized("平台SKU")]).fill = yellow_fill
        exclusions.add(platform_sku)
        appended += 1
    return appended, len(legacy_platform_skus)


def platform_sales_eligible_rows(sheet, cols):
    """Return rows selected by the original Step2 mapping condition."""
    result = []
    for row in range(3, sheet.max_row + 1):
        platform_sku = clean(sheet.cell(row, cols[normalized("平台SKU")]).value)
        warehouse_sku = clean(sheet.cell(row, cols[normalized("仓库SKU")]).value)
        if (
            platform_sku
            and warehouse_sku
            and not any(
                clean(sheet.cell(row, cols[normalized(name)]).value)
                for name in BASE_FIELDS
            )
        ):
            result.append(row)
    return result


def site_rank(alias):
    text = clean(alias).upper()
    if text.endswith("-US"):
        return 0
    if text.endswith("-CA"):
        return 1
    return 2


def load_platform(path):
    source = pd.read_excel(path, sheet_name="平台SKU映射&备货", header=1, dtype=str).fillna("")
    require({normalized(c): c for c in source.columns}, ["主体", "平台SKU", "系统SKU", "店铺别名", "运营", "备货状态", "小组"], "平台SKU映射")
    canonical = source["主体"].map(canonical_subject)
    outside = source.loc[canonical == "", "主体"].map(clean)
    outside = sorted({value for value in outside if value})
    if outside:
        raise ValueError(f"平台SKU映射包含目标范围外主体: {', '.join(outside[:20])}")
    return source, source.copy()


def load_sales(path):
    columns = ["店铺", "MSKU", "ASIN", "小组"] + [f"标签{i:02d}" for i in range(1, 11)]
    source = pd.read_excel(path, sheet_name="MSKU销售面板", dtype=str).fillna("")
    require({normalized(c): c for c in source.columns}, columns, "亚马逊销售面板")
    matching = source.loc[group_mask(source["小组"]), columns].copy()
    return source, matching


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--tracking", required=True, type=Path)
    parser.add_argument("--products", required=True, type=Path)
    parser.add_argument("--platform", type=Path)
    parser.add_argument("--sales", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--fresh-after", required=True, type=float, help="Unix timestamp; every ERP export must be downloaded after this run start")
    return parser


def main():
    args = build_parser().parse_args()
    if bool(args.platform) != bool(args.sales):
        raise ValueError("--platform 和 --sales 必须同时提供或同时省略")
    inputs = [args.workbook, args.tracking, args.products] + ([args.platform, args.sales] if args.platform else [])
    if any(not path.is_file() for path in inputs):
        raise ValueError("一个或多个输入文件不存在")
    if args.output.resolve() in {path.resolve() for path in inputs}:
        raise ValueError("输出文件不能覆盖任何输入文件")
    if args.output.exists():
        raise ValueError(f"输出文件已存在: {args.output}")
    stale = [str(path) for path in inputs[1:] if path.stat().st_mtime < args.fresh_after]
    if stale:
        raise ValueError("检测到非本次运行下载的 ERP 输入文件，已拒绝使用: " + ", ".join(stale))

    tracking_source = pd.read_excel(args.tracking, dtype=str).fillna("")
    tracking = prepare_tracking_data(tracking_source)

    product_fields = ["产品SKU", "产品名称", "三级品类", "品类等级", "产品销售状态", "产品实际等级", "产品目标等级"]
    products_source = pd.read_excel(args.products, dtype=str).fillna("")
    require({normalized(c): c for c in products_source.columns}, product_fields, "产品档案")
    products = products_source.loc[:, product_fields].copy()
    product_map = products.drop_duplicates("产品SKU", keep="first").set_index("产品SKU")
    tracking_map = tracking.set_index("平台SKU")
    if args.platform:
        platform_source, platform = load_platform(args.platform)
        sales_source, sales = load_sales(args.sales)
    else:
        platform_source = platform = pd.DataFrame(columns=["主体", "平台SKU", "系统SKU", "店铺别名", "运营", "备货状态", "小组"])
        sales_source = sales = pd.DataFrame(columns=["店铺", "MSKU", "ASIN", "小组"] + [f"标签{i:02d}" for i in range(1, 11)])

    workbook = load_workbook(args.workbook)
    if "产品表" not in workbook.sheetnames:
        raise ValueError("工作簿缺少产品表")
    legacy_name = next((name for name in LEGACY_NAMES if name in workbook.sheetnames), None)
    legacy_before = sheet_digest(workbook[legacy_name]) if legacy_name else None
    preexisting_errors = formula_errors(workbook)

    sheet = workbook["产品表"]
    cols = headers(sheet)
    require(cols, REQUIRED_MAIN, "产品表")
    appended, legacy_sku_count = append_tracking_skus(
        workbook, sheet, cols, tracking, legacy_name
    )

    tracking_targets = {"新品期开售月": "预计开售时间", "新品期截止月": "新品期截止月", "最新新品成功率": "预测销量完成比例"}
    product_targets = {"产品名称": "产品名称", "三级品类": "三级品类", "产品销售状态": "产品销售状态", "品类等级": "产品实际等级"}
    tracking_matches = product_matches = 0
    for row in range(3, sheet.max_row + 1):
        platform_sku = clean(sheet.cell(row, cols[normalized("平台SKU")]).value)
        warehouse_sku = clean(sheet.cell(row, cols[normalized("仓库SKU")]).value)
        if platform_sku in tracking_map.index:
            tracking_matches += 1
            source = tracking_map.loc[platform_sku]
            for target, source_name in tracking_targets.items():
                sheet.cell(row, cols[normalized(target)]).value = clean(source[source_name]) or None
        if warehouse_sku in product_map.index:
            product_matches += 1
            source = product_map.loc[warehouse_sku]
            for target, source_name in product_targets.items():
                sheet.cell(row, cols[normalized(target)]).value = clean(source[source_name]) or None

    platform_map = {}
    for _, source in platform.iterrows():
        sku, alias = clean(source["平台SKU"]), clean(source["店铺别名"])
        if not sku or not alias or "老三" in alias:
            continue
        candidate = {"店铺别名": alias, "主体": product_subject(source["主体"]), "运营": clean(source["运营"]), "备货状态": clean(source["备货状态"]), "rank": site_rank(alias)}
        if sku not in platform_map or candidate["rank"] < platform_map[sku]["rank"]:
            platform_map[sku] = candidate
    sales_map = {}
    for _, source in sales.iterrows():
        key = (clean(source["店铺"]), clean(source["MSKU"]))
        if all(key) and key not in sales_map:
            sales_map[key] = {"ASIN": clean(source["ASIN"]), **{f"标签{i:02d}": clean(source[f"标签{i:02d}"]) for i in range(1, 11)}}

    eligible_rows = platform_sales_eligible_rows(sheet, cols)
    if not args.platform and eligible_rows:
        raise ValueError(
            "旧 Step2 逻辑检测到需要补充平台 SKU 映射的行，"
            "必须同时提供 --platform 和 --sales"
        )

    step2_eligible = len(eligible_rows)
    step2_platform = step2_sales = 0
    for row in eligible_rows:
        platform_sku = clean(sheet.cell(row, cols[normalized("平台SKU")]).value)
        source = platform_map.get(platform_sku)
        if not source:
            continue
        step2_platform += 1
        for name in BASE_FIELDS:
            sheet.cell(row, cols[normalized(name)]).value = source[name] or None
        sale = sales_map.get((source["店铺别名"], platform_sku))
        if sale:
            step2_sales += 1
            sheet.cell(row, cols[normalized("ASIN")]).value = sale["ASIN"] or None
            for index in range(1, 11):
                name = f"标签{index:02d}"
                sheet.cell(row, cols[normalized(name)]).value = sale[name] or None

    subjects_normalized = normalize_product_subjects(sheet, cols)

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bordered = 0
    for row in range(3, sheet.max_row + 1):
        if any(clean(sheet.cell(row, cols[normalized(key)]).value) for key in ("平台SKU", "仓库SKU")):
            bordered += 1
            for column in range(1, 44):
                sheet.cell(row, column).border = border

    replace_sheet(workbook, "ERP新品成功率", tracking_source)
    replace_sheet(workbook, "ERP产品档案", products_source)
    if args.platform:
        replace_sheet(workbook, "Step2平台SKU映射", platform_source)
        replace_sheet(workbook, "Step2亚马逊销售面板", sales_source)

    legacy_after = sheet_digest(workbook[legacy_name]) if legacy_name else None
    if legacy_name and legacy_before != legacy_after:
        raise AssertionError("产品表-老三(含清货淘汰)发生变化，已停止保存")

    current_errors = formula_errors(workbook)
    new_errors = sorted(current_errors - preexisting_errors)
    if new_errors:
        raise AssertionError("更新新增了公式错误: " + "; ".join(new_errors[:20]))

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    saved = load_workbook(args.output)
    remaining_short_subjects = short_product_subject_cells(
        saved["产品表"], headers(saved["产品表"])
    )
    if remaining_short_subjects:
        saved.close()
        args.output.unlink(missing_ok=True)
        raise AssertionError(
            "保存后产品表仍含短格式主体: "
            + ", ".join(remaining_short_subjects[:20])
        )
    if legacy_name and sheet_digest(saved[legacy_name]) != legacy_before:
        args.output.unlink(missing_ok=True)
        raise AssertionError("保存后产品表-老三(含清货淘汰)摘要变化，输出已删除")
    saved.close()
    print(json.dumps({
        "output": str(args.output.resolve()),
        "trackingSourceRows": len(tracking_source),
        "trackingMatchingRows": len(tracking),
        "productSourceRows": len(products_source),
        "platformSourceRows": len(platform_source) if args.platform else None,
        "salesSourceRows": len(sales_source) if args.sales else None,
        "trackingMatches": tracking_matches,
        "productMatches": product_matches,
        "newSkuAppended": appended,
        "legacySkuExclusions": legacy_sku_count,
        "legacySheetPresent": bool(legacy_name),
        "step2Eligible": step2_eligible,
        "step2PlatformMatches": step2_platform,
        "step2SalesMatches": step2_sales,
        "subjectsNormalizedToLongFormat": subjects_normalized,
        "platformSalesSourcesRefreshed": bool(args.platform),
        "borderedRows": bordered,
        "legacySheetUnchanged": True if legacy_name else None,
        "preexistingFormulaErrors": len(preexisting_errors),
        "newFormulaErrors": 0,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
